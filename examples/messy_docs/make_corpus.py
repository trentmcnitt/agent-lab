"""Writes the messy folder: what a support team's shared drive looks like after a few years.

    uv run python make_corpus.py            # (re)writes ./docs, byte-for-byte the same every time

Mixed formats (.md .txt .pdf .html .docx), names nobody chose carefully ("FINAL_v2", "Copy of",
dates, spaces, a sub-folder or two), an exact duplicate and a stale near-duplicate, files with
no headings, a PDF whose metadata title is reportlab's "untitled", a 24-page handbook whose
metadata title is "Microsoft Word - …", and a file that is nearly empty. Nothing here was
written for Agent Lab: the point is to see what the bench can show from files like these.
"""
from __future__ import annotations

import random
import shutil
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
DOCS = HERE / "docs"

REFUND_V2 = """# Refund Policy

Last reviewed March 2024 by the Customer Care leads.

## Who can get a refund

Any customer who bought a Kestrel appliance directly from kestrelkitchen.com or from an authorized retailer can ask for a refund. Gift recipients can ask too, but the refund goes back to the original payment method unless the buyer agrees otherwise.

## Time limits

Customers have 45 days from the delivery date to return an item for a full refund. After 45 days and up to 90 days we offer store credit only. Past 90 days a return is not possible, but the warranty may still apply (see the warranty terms).

Opened small appliances are fine to return as long as all parts are included. Blades, filters and gaskets that have been used are not returnable on their own.

## How to process it

1. Confirm the order number and delivery date in the order system.
2. Send the customer a prepaid return label from the Returns tab.
3. Once the warehouse scans the box, issue the refund. Do not refund before the scan unless a lead approves it.
4. Refunds show up on the customer's card in 5 to 10 business days.

## Exceptions

Leads can approve a refund outside the window for a defective unit, a shipping error on our side, or a customer who was quoted the wrong policy by an agent. Note the reason in the ticket.
"""

REFUND_OLD = REFUND_V2.replace("45 days", "30 days").replace("March 2024", "June 2022").replace(
    "5 to 10 business days", "7 to 14 business days")

SHIPPING_NOTES = """called warehouse again re: the Ohio delays. they say the regional carrier is backed up til end of month. for now tell customers 7-10 business days for standard shipping to OH, PA, WV. expedited still 2-3 days.

if a package shows delivered but customer says they don't have it, wait 48 hours then open a carrier trace. don't reship before the trace comes back unless its a lead approval.

PO boxes: standard only, no expedited. APO/FPO fine with standard.

reminder from Dana: damaged-in-transit claims need photos of the box AND the unit. no photos = no claim with the carrier, but we can still replace as a goodwill thing if under $80.
"""

SHIPPING_RATES = """Shipping rates 2024
Standard (5-7 business days): free on orders over $49, otherwise $6.95
Expedited (2-3 business days): $14.95
Overnight (order by 1pm ET): $29.95
Alaska and Hawaii: standard only, add $12
Canada: flat $19.95, duties paid by the customer at delivery
We do not ship to other countries yet.
"""

WARRANTY_PAGES = [
    """Kestrel Kitchen Limited Warranty

This limited warranty covers Kestrel blenders, mixers, toasters and kettles bought new from Kestrel Kitchen or an authorized retailer. It gives the original owner specific legal rights, and they may also have other rights that vary from state to state.

What is covered in the first year: any defect in materials or workmanship. We will repair or replace the unit at no cost, including shipping both ways.""",
    """Years two through five: the motor base only. If the motor fails under normal household use, we replace the motor base. Jars, lids, blades, gaskets and other wear parts are not covered after the first year. The customer pays shipping to us; we pay shipping back.

What is never covered: commercial use, damage from dropping or misuse, units bought from unauthorized sellers, and cosmetic wear.""",
    """How to make a claim: the customer needs proof of purchase and the serial number from the bottom of the motor base. Agents open a warranty claim in the order system and attach both. Replacement units ship within 3 business days of an approved claim.

Kestrel Kitchen, 1200 Harbor Way, Portland OR. Warranty document KK-W-2023.""",
]

LOANER_PAGES = [
    """LOANER UNIT PROGRAM

Effective immediately we can offer a loaner blender while a customer's unit is in for warranty repair, if the repair is expected to take more than 10 business days. Loaners are the Model B200 only. Ask a lead to release one from the loaner pool.

The customer returns the loaner in the same box their repaired unit arrives in. If the loaner is not returned within 14 days of their repaired unit being delivered, charge the card on file $89.""",
    """Loaners are not offered for mixers, toasters or kettles. Loaners are not offered outside the United States.

Questions: Customer Care leads channel.""",
]

CHARGEBACK_PAGES = [
    """When a customer disputes a charge with their bank, the bank tells us through the payment processor. Do not contact the bank. Gather the order record, the delivery confirmation and any ticket history, and send them to billing@ within 5 business days.

If the customer contacts us after filing a dispute, we cannot also issue a refund: the bank's decision settles it. Explain this politely.""",
    """Chargeback fees ($15 per dispute) are charged to us by the processor, not to the customer. Never pass this fee on to a customer and never offer to reimburse a customer for a bank fee they were charged.""",
]

HANDBOOK_TOPICS = [
    ("Welcome to Customer Care", "our team, how we work, and what a good day looks like"),
    ("Your first week", "shadowing, system access, and the certification quiz"),
    ("Tone and empathy", "how we write to customers and what to avoid"),
    ("The order system", "finding orders, notes, and the returns tab"),
    ("Phone etiquette", "greetings, holds, and transfers"),
    ("Chat and email", "response times, templates, and signatures"),
    ("Refunds and store credit", "who can approve what, and when a lead is needed"),
    ("Warranty claims", "proof of purchase, serial numbers, and replacement units"),
    ("Shipping problems", "lost, late, and damaged packages"),
    ("Product knowledge: blenders", "models, speeds, and common questions"),
    ("Product knowledge: mixers", "attachments, bowls, and speeds"),
    ("Product knowledge: kettles and toasters", "settings and safety"),
    ("Difficult conversations", "angry customers, threats, and when to hand off"),
    ("Escalations", "what goes to a lead and what goes to legal"),
    ("Privacy and security", "verifying identity and what never to say on a call"),
    ("Quality reviews", "how tickets are scored each month"),
    ("Schedules and time off", "shift swaps, holidays, and sick days"),
    ("Tools we use", "the order system, the phone system, and the knowledge base"),
    ("Metrics", "first-response time, resolution time, and satisfaction"),
    ("Holiday season", "the November and December playbook"),
    ("Accessibility", "helping customers who use assistive technology"),
    ("Recalls", "what to do if a product is recalled"),
    ("Feedback to product", "how customer feedback reaches the product team"),
    ("Leaving the team", "offboarding and returning equipment"),
]

SENTENCES = [
    "New agents should read this section before their first live shift.",
    "When in doubt, ask in the leads channel rather than guessing.",
    "Write down what you promised the customer in the ticket so the next agent can follow through.",
    "We never argue with a customer about what they experienced; we work out what we can do.",
    "If a policy and this handbook disagree, the policy document wins and you should tell a lead.",
    "Every exception you make should have a reason in the ticket notes.",
    "Use the customer's name, and keep sentences short.",
    "Our target is a first response within four business hours on email and two minutes on chat.",
    "Avoid jargon such as SKU or RMA when talking to customers.",
    "Supervisors review a sample of tickets from every agent each month.",
    "Remember that most customers contact us only when something has gone wrong.",
    "A calm, specific answer is better than a fast vague one.",
]


def _handbook_pages() -> list[str]:
    rng = random.Random(7)
    pages = []
    for title, about in HANDBOOK_TOPICS:
        body = [f"{title}", f"This section covers {about}."]
        for _ in range(9):
            body.append(" ".join(rng.sample(SENTENCES, 4)))
        pages.append("\n\n".join(body))
    return pages


def _pdf(path: Path, pages: list[str], title: str | None) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Frame, Paragraph

    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)   # invariant: same bytes every time
    if title is not None:
        c.setTitle(title)                                          # else reportlab writes "untitled"
    style = getSampleStyleSheet()["BodyText"]
    for page in pages:
        frame = Frame(72, 72, letter[0] - 144, letter[1] - 144)
        frame.addFromList([Paragraph(p.replace("\n", " "), style) for p in page.split("\n\n")], c)
        c.showPage()
    c.save()


def _docx(path: Path, blocks: list[tuple[str, str]]) -> None:
    import docx

    d = docx.Document()
    for kind, text in blocks:
        if kind == "h":
            d.add_heading(text, level=1)
        else:
            d.add_paragraph(text)
    d.core_properties.title = ""
    d.save(str(path))
    _fix_zip_dates(path)


def _fix_zip_dates(path: Path) -> None:
    """python-docx stamps the zip entries with today's time; rewrite them so the file is stable."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re
                data = re.sub(rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>", b"", data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2024, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, data)
    tmp.replace(path)


FILES: dict[str, str] = {
    "Refund Policy FINAL_v2.md": REFUND_V2,
    "Copy of Refund Policy FINAL_v2.md": REFUND_V2,                     # exact duplicate
    "refund policy (old).md": REFUND_OLD,                               # stale near-duplicate
    "2023-01-15 shipping notes.txt": SHIPPING_NOTES,                    # no headings, a date name
    "shipping_rates_2024.txt": SHIPPING_RATES,
    "notes.txt": "TODO ask Dana about the new label printer\n",          # nearly empty
    "untitled.md": "\n",                                                # empty
    "faq.md": """# Frequently Asked Questions

## 1. Where is my order?
Orders ship within 2 business days. The tracking link is in the shipping confirmation email.

## 2. Can I change my order after placing it?
Within 1 hour of ordering, yes: cancel and reorder. After that the warehouse may already have it.

## 3. Do you price match?
We match our own site's price within 14 days of purchase. We do not match other retailers.

## 4. Is the blender jar dishwasher safe?
Yes, top rack only. The blade assembly should be hand-washed.

## 5. How do I register my product?
At kestrelkitchen.com/register with the serial number from the bottom of the base.
""",
    "billing/Billing FAQ.md": """# Billing FAQ

## Why was I charged twice?
Usually one charge is a pending authorization that falls off in 3 to 5 days. If both post, refund the duplicate.

## Can I pay in installments?
Orders over $150 can use the installment option at checkout. Agents cannot set up installments after the order is placed.

## Sales tax
We collect sales tax in every state that requires it. Tax-exempt customers send their certificate to billing@ before ordering.
""",
    "billing/invoice-disputes.txt": """Invoice disputes (business accounts only)

Business customers who think an invoice is wrong email billing@ with the invoice number. Billing answers within 3 business days. Agents should not adjust business invoices themselves.
""",
    "onboarding/new agent checklist.md": """# New agent checklist

- [ ] Order system login
- [ ] Phone system login and voicemail PIN
- [ ] Shadow two shifts with a senior agent
- [ ] Read the Refund Policy and the Warranty terms
- [ ] Pass the certification quiz (80% or better)
- [ ] First live shift with a lead on the channel
""",
    "onboarding/Tone & Voice guide.md": """We sound like a helpful neighbor who happens to know a lot about kitchen appliances. Warm, direct, never cute.

Say "I can do that for you" instead of "Unfortunately our policy states". Say sorry once, then fix it. Don't use exclamation marks more than once per message.

Sign off with your first name only.
""",
    "troubleshooting/Device won't turn on.md": """## Blender or mixer won't turn on

1. Check that the jar or bowl is locked onto the base. The safety switch keeps the motor off until it clicks.
2. Try a different outlet. Kitchen outlets on a tripped GFCI are the most common cause.
3. Hold the power button for 5 seconds to reset the controller.
4. If the light blinks red three times, the motor overheated: unplug it for 30 minutes.

If none of this works, it is a warranty claim.
""",
    "troubleshooting/wifi pairing steps.txt": """Smart kettle wifi pairing
hold the temp button + power for 6 sec until the light flashes blue
open the Kestrel app > add device > kettle
2.4 GHz networks only, 5 GHz will not work
if it fails 3 times, factory reset: hold power 15 sec
""",
    "SLA.md": """# Service levels

Email: first response within 4 business hours. Chat: within 2 minutes during staffed hours. Phone: 80% of calls answered within 60 seconds.

Staffed hours are 8am to 8pm ET Monday to Friday and 10am to 4pm ET Saturday.
""",
    "holiday hours 2023.txt": """Thanksgiving: closed
Black Friday: 7am-10pm ET (all hands)
Dec 24: 8am-2pm ET
Dec 25: closed
Dec 31: 8am-4pm ET
Jan 1: closed
""",
    "macros.txt": """--- MACRO: refund approved ---
Hi {name}, good news: your refund is approved. You'll see it on your card in 5 to 10 business days. Thanks for your patience, {agent}

--- MACRO: need serial number ---
Hi {name}, to get your warranty claim started I need the serial number from the bottom of the motor base. A photo works great. {agent}

--- MACRO: shipping delay ---
Hi {name}, our carrier is running a few days behind in your area. Your order is on its way and the tracking link will update as it moves. {agent}
""",
    "kb-article-1187.html": """<html><body>
<h1>Kettle shows E2 or won't heat</h1>
<p>E2 means the kettle sensed it was empty. Fill above the MIN line and press start again.</p>
<p>If the kettle won't heat at all and shows no code, check the base is seated and the outlet works.</p>
</body></html>
""",
    "password reset.html": """<html><head><title>Resetting a customer account password</title></head><body>
<p>Agents can't see or set passwords. Send the customer the reset link from the account page: Account, then Security, then Send reset email.</p>
<p>The link expires after 30 minutes. If the email doesn't arrive, check the address on file and the customer's spam folder.</p>
</body></html>
""",
    "troubleshooting/error codes.html": """<html><head><title>Blender error codes</title></head><body>
<table>
<tr><th>Code</th><th>Meaning</th><th>What to tell the customer</th></tr>
<tr><td>E1</td><td>Jar not detected</td><td>Lock the jar onto the base until it clicks.</td></tr>
<tr><td>E3</td><td>Blade jammed</td><td>Unplug, remove the jar, and clear food from around the blades.</td></tr>
<tr><td>E4</td><td>Motor overheated</td><td>Unplug for 30 minutes and let it cool; blend in shorter bursts with more liquid.</td></tr>
<tr><td>E7</td><td>Controller fault</td><td>Hold power for 5 seconds; if E7 returns, open a warranty claim.</td></tr>
</table>
</body></html>
""",
}


def build(out: Path = DOCS) -> Path:
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    for rel, text in FILES.items():
        p = out / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text)
    _pdf(out / "Warranty Terms.pdf", WARRANTY_PAGES, title="Kestrel Kitchen Limited Warranty")
    _pdf(out / "scan_0042.pdf", LOANER_PAGES, title=None)                   # metadata title "untitled"
    _pdf(out / "billing/chargebacks_2021.pdf", CHARGEBACK_PAGES, title="Chargebacks & Disputes: Agent Guide")
    _pdf(out / "Employee Handbook - Customer Care (2022).pdf", _handbook_pages(),
         title="Microsoft Word - CC_Handbook_2022_rev3.docx")                # junk metadata title
    _docx(out / "Returns process.docx", [
        ("h", "Returns process"),
        ("p", "A return starts when the customer asks for one in the window described in the Refund Policy."),
        ("p", "Create the return in the Returns tab, send the prepaid label, and set the ticket to Waiting on customer."),
        ("p", "When the warehouse scans the package the ticket reopens automatically; issue the refund then."),
    ])
    _docx(out / "escalation matrix v3 FINAL FINAL.docx", [
        ("p", "Escalate to a lead: refunds outside the window, goodwill over $80, any customer who mentions a lawyer."),
        ("p", "Escalate to legal: injury reports, property damage, any recall question."),
        ("p", "Escalate to billing: chargebacks, business invoices, tax-exempt certificates."),
    ])
    return out


if __name__ == "__main__":
    print(build())
